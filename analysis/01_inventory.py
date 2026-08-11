#!/usr/bin/env python3
"""Create a non-analytic inventory of the locked SAP and raw NEDIS workbook."""

from __future__ import annotations

import hashlib
import json
import platform
import sys
from pathlib import Path

import openpyxl

from settings import RAW_XLSX, ROOT, SAP_DOCX

OUT = ROOT / "analysis_outputs" / "00_inventory"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def display_path(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def main() -> None:
    if not RAW_XLSX.is_file():
        raise FileNotFoundError(
            f"Restricted NEDIS workbook not found: {RAW_XLSX}. "
            "Set NEDIS_RAW_XLSX or follow data/README.md."
        )
    OUT.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(RAW_XLSX, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        sheets.append(
            {
                "name": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "header": [None if x is None else str(x) for x in first_row],
            }
        )

    variable_ws = wb["변수설명"]
    variable_rows = [
        [None if x is None else str(x) for x in row]
        for row in variable_ws.iter_rows(values_only=True)
    ]
    wb.close()

    sap_inventory = {
        "path": display_path(SAP_DOCX),
        "present": SAP_DOCX.is_file(),
        "lock_date_rendered": "2026-08-04",
    }
    if SAP_DOCX.is_file():
        sap_inventory.update(
            {
                "size_bytes": SAP_DOCX.stat().st_size,
                "sha256": sha256(SAP_DOCX),
            }
        )

    inventory = {
        "created_utc": __import__("datetime").datetime.now(__import__("datetime").timezone.utc).isoformat(),
        "python": sys.version,
        "platform": platform.platform(),
        "raw_workbook": {
            "path": display_path(RAW_XLSX),
            "size_bytes": RAW_XLSX.stat().st_size,
            "sha256": sha256(RAW_XLSX),
        },
        "locked_sap": sap_inventory,
        "worksheets": sheets,
    }
    (OUT / "inventory.json").write_text(
        json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUT / "variable_description_rows.json").write_text(
        json.dumps(variable_rows, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(inventory, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
